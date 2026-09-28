import os
import telebot
from google import genai
from docx import Document
from PIL import Image
import json

# Получаем токены из переменных окружения сервера (безопасный метод для GitHub и Render)
TELEGRAM_TOKEN = os.environ.get("TELEGRAM_TOKEN", "ВАШ_ТОКЕН_ТЕЛЕГРАМА")
GEMINI_API_KEY = os.environ.get("GEMINI_API_KEY", "ВАШ_КЛЮЧ_GEMINI")

bot = telebot.TeleBot(TELEGRAM_TOKEN)
client = genai.Client(api_key=GEMINI_API_KEY)
TEMPLATE_PATH = "Бланк автоматической загрузки.docx"

@bot.message_handler(commands=['start'])
def send_welcome(message):
    bot.reply_to(message, "Бот успешно запущен на сервере! Пришлите фото заполненного бланка.")

@bot.message_handler(content_types=['photo'])
def handle_photo(message):
    bot.reply_to(message, "⏳ Распознаю почерк через серверы Google, подождите пару секунд...")
    
    image_path = "temp_blank.jpg"
    fixed_image_path = "fixed_blank.jpg"
    try:
        # 1. Скачиваем фото из Telegram
        file_info = bot.get_file(message.photo[-1].file_id)
        downloaded_file = bot.download_file(file_info.file_path)
        with open(image_path, 'wb') as new_file:
            new_file.write(downloaded_file)

        # 2. Принудительно конвертируем в чистый JPG (решает проблему с ошибками форматов)
        img = Image.open(image_path)
        if img.mode != 'RGB':
            img = img.convert('RGB')
        img.save(fixed_image_path, "JPEG")
        
        # 3. Открываем исправленное фото для отправки в Google
        image_for_gemini = Image.open(fixed_image_path)

        prompt = """
        Внимательно посмотри на этот рукописный бланк заказа. Извлеки все заполненные от руки данные.
        Верни результат СТРОГО в формате JSON. Ключи должны точно соответствовать полям на бланке: 
        номер_заказа, Дата_изготовления, ФИО_умершего, Дата_рождения, Дата_смерти, Обелиск, Тумба, Цветник_Надгробная, Портрет, Крест, Цветы, Свеча, Эпитафия, Сумма, Установка, Предоплата, Итог, Остаток, Адрес_кладбища, Телефон, ФИО_заказчика, Кем_принят_заказ, Дата_приема.
        Если какое-то поле не заполнено, оставь пустую строку "". Верни ТОЛЬКО чистый JSON, без лишнего текста.
        """

        # 4. Отправляем запрос к мощной модели Gemini
        response = client.models.generate_content(
            model="gemini-2.5-flash",
            contents=[prompt, image_for_gemini]
        )
        
        text_response = response.text.strip()
        print("\n--- РАСПОЗНАНО GOOGLE GEMINI ---")
        print(text_response)
        
        # 5. Очищаем ответ от лишнего текста, оставляя только JSON
        if "{" in text_response and "}" in text_response:
            start = text_response.find("{")
            end = text_response.rfind("}") + 1
            text_response = text_response[start:end]
            
        data = json.loads(text_response)

        # 6. Вставляем данные в Word-шаблон
        doc = Document(TEMPLATE_PATH)
        for key, value in data.items():
            if not value:
                continue
            target_tags = [f"{{{{{key}}}}}", f"{{{key}}}", key]
            
            for paragraph in doc.paragraphs:
                for tag in target_tags:
                    if tag in paragraph.text:
                        paragraph.text = paragraph.text.replace(tag, str(value))
                        
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for tag in target_tags:
                            if tag in cell.text:
                                cell.text = cell.text.replace(tag, str(value))

        output_filename = "Готовый_бланк.docx"
        doc.save(output_filename)

        # 7. Отправляем документ клиенту
        with open(output_filename, 'rb') as doc_file:
            bot.send_document(message.chat.id, doc_file, caption="✅ Готово! Бланк заполнен.")

    except Exception as e:
        bot.reply_to(message, f"❌ Произошла ошибка: {e}")
        print(f"ОШИБКА: {e}")
    finally:
        # Удаляем временные файлы, чтобы не засорять сервер
        if os.path.exists(image_path):
            os.remove(image_path)
        if os.path.exists(fixed_image_path):
            os.remove(fixed_image_path)

bot.infinity_polling()
